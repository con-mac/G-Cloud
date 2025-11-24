"""
Document generation service for G-Cloud proposals (PA Deployment)
Uses SharePoint Online instead of Azure Blob Storage
PLACEHOLDER: SharePoint integration needs to be completed
"""

import os
import logging
from pathlib import Path
from typing import Dict, List, Optional
from docx import Document
import uuid

logger = logging.getLogger(__name__)


class DocumentGenerator:
    """Generates G-Cloud proposal documents from templates"""
    
    def __init__(self):
        """Initialize document generator for PA deployment (SharePoint)"""
        # Check if we're using SharePoint
        self.use_sharepoint = os.environ.get("USE_SHAREPOINT", "false").lower() == "true"
        
        # PLACEHOLDER: Initialize SharePoint service
        # from sharepoint_service.sharepoint_online import upload_file_to_sharepoint
        # self.sharepoint_service = get_sharepoint_service()
        
        # Use /tmp for temporary files (Azure Functions)
        self.templates_dir = Path("/tmp/templates")
        self.output_dir = Path("/tmp/generated_documents")
        self.output_dir.mkdir(parents=True, exist_ok=True)
    
    def generate_service_description(
        self,
        title: str,
        description: str,
        features: List[str],
        benefits: List[str],
        service_definition: List[dict] | None = None,
        update_metadata: Dict | None = None,
        save_as_draft: bool = False,
        new_proposal_metadata: Dict | None = None
    ) -> Dict[str, str]:
        """
        Generate Service Description document from template
        PLACEHOLDER: SharePoint upload needs to be implemented
        
        Returns:
            Dict with SharePoint item IDs or local paths
        """
        logger.warning("generate_service_description: Using placeholder implementation")
        
        # PLACEHOLDER: Load template from SharePoint or local
        # For now, use local template path
        template_path = self.templates_dir / "service_description_template.docx"
        
        if not template_path.exists():
            # Try alternative paths
            alt_paths = [
                Path("/app/templates/service_description_template.docx"),
                Path("/app/docs/service_description_template.docx"),
                Path("docs/service_description_template.docx"),
            ]
            for alt_path in alt_paths:
                if alt_path.exists():
                    template_path = alt_path
                    break
        
        if not template_path.exists():
            raise FileNotFoundError(f"Template not found: {template_path}")
        
        # Load template
        doc = Document(template_path)
        
        # PLACEHOLDER: Replace placeholders in document
        # This is a simplified version - full implementation needed
        for paragraph in doc.paragraphs:
            if 'ENTER SERVICE NAME HERE' in paragraph.text:
                paragraph.text = paragraph.text.replace('ENTER SERVICE NAME HERE', title)
            if '{{SERVICE_NAME}}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{{SERVICE_NAME}}', title)
        
        # Generate filename
        if update_metadata or new_proposal_metadata:
            gcloud_version = (update_metadata or new_proposal_metadata).get('gcloud_version', '15')
            service_name = (update_metadata or new_proposal_metadata).get('service_name') or (new_proposal_metadata or {}).get('service', title)
            word_filename = f"PA GC{gcloud_version} SERVICE DESC {service_name}.docx"
            
            if save_as_draft:
                word_filename = word_filename.replace('.docx', '_draft.docx')
        else:
            doc_id = str(uuid.uuid4())[:8]
            word_filename = f"{title}_{doc_id}.docx"
        
        # Save to temp location
        word_path = self.output_dir / word_filename
        doc.save(str(word_path))
        
        # PLACEHOLDER: Upload to SharePoint
        if self.use_sharepoint and (update_metadata or new_proposal_metadata):
            # PLACEHOLDER: Upload to SharePoint using Graph API
            # sharepoint_item_id = upload_file_to_sharepoint(word_path, target_path, gcloud_version)
            logger.warning("SharePoint upload not yet implemented - file saved locally only")
            sharepoint_item_id = None
        else:
            sharepoint_item_id = None
        
        return {
            "word_path": str(word_path),
            "word_filename": word_filename,
            "word_sharepoint_id": sharepoint_item_id,  # PLACEHOLDER
            "pdf_path": None,  # PLACEHOLDER: PDF conversion
            "pdf_filename": None,
            "pdf_sharepoint_id": None,  # PLACEHOLDER
        }

